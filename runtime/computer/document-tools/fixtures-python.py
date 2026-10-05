import gzip
import json
import os
import pathlib
import subprocess
import sys
import zipfile

import pandas
import pdfplumber
import pymupdf
import pypdf
import pytesseract
import xlsxwriter
from bs4 import BeautifulSoup
from defusedxml import ElementTree
from docx import Document
from ebooklib import epub
from lxml import etree
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas
import yaml


root, work = map(pathlib.Path, sys.argv[1:3])
marker = "OPULENT DOCUMENT FIXTURE"
tessdata = "/usr/share/tesseract-ocr/5/tessdata"
env = {"PATH": "/usr/local/bin:/usr/bin:/bin", "HOME": str(work), "LANG": "C.UTF-8",
       "TESSDATA_PREFIX": tessdata}
proof = {}


def run(name, args):
    result = subprocess.run([name, *map(str, args)], check=True, capture_output=True,
                            text=True, timeout=60, env=env)
    return result.stdout


def contains(text, name):
    if marker not in text:
        raise RuntimeError(name + " failed to recover the real fixture text")
    proof[name] = True


pdf = work / "fixture.pdf"
writer = canvas.Canvas(str(pdf))
writer.drawString(72, 720, marker)
writer.drawString(72, 690, "Article 1 - local document extraction")
writer.save()
with pymupdf.open(pdf) as document:
    contains(document[0].get_text(), "pymupdf-text")
    document[0].get_pixmap(dpi=100).save(work / "pdf-render.png")
    if not (work / "pdf-render.png").stat().st_size:
        raise RuntimeError("PyMuPDF returned an empty render")
proof["pymupdf-render"] = True
contains(pypdf.PdfReader(pdf).pages[0].extract_text(), "pypdf-text")
with pdfplumber.open(pdf) as document:
    contains(document.pages[0].extract_text(), "pdfplumber-text")
docx = work / "fixture.docx"
document = Document()
document.add_heading(marker, 0)
document.add_paragraph(marker)
document.add_paragraph("The archive contains local evidence, not a hosted provider response.")
table = document.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Article"
table.cell(0, 1).text = "Value"
table.cell(1, 0).text = "1"
table.cell(1, 1).text = "42"
document.save(docx)
contains("\n".join(p.text for p in Document(docx).paragraphs), "python-docx-roundtrip")
book = Workbook()
book.active["A1"] = marker
book.active["A2"] = 42
book.save(work / "fixture.xlsx")
contains(load_workbook(work / "fixture.xlsx").active["A1"].value, "openpyxl-roundtrip")
with xlsxwriter.Workbook(work / "xlsxwriter.xlsx") as sheet:
    sheet.add_worksheet().write(0, 0, marker)
contains(load_workbook(work / "xlsxwriter.xlsx").active["A1"].value, "xlsxwriter-roundtrip")
slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[6])
slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text = marker
slides.save(work / "fixture.pptx")
contains("\n".join(shape.text for shape in Presentation(work / "fixture.pptx").slides[0].shapes if shape.has_text_frame), "python-pptx-roundtrip")
image = Image.new("RGB", (1800, 500), "white")
ImageDraw.Draw(image).text((80, 140), marker,
                          font=ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 72), fill="black")
image.save(work / "scan.png")
scanned = pymupdf.open()
page = scanned.new_page(width=900, height=250)
page.insert_image(page.rect, filename=str(work / "scan.png"))
scanned.save(work / "scan.pdf")
scanned.close()
contains(run("tesseract", [work / "scan.png", "stdout", "-l", "eng", "--tessdata-dir", tessdata]), "tesseract-offline-ocr")
contains(pytesseract.image_to_string(image, lang="eng", config="--tessdata-dir " + tessdata), "pytesseract-offline-ocr")
with pymupdf.open(work / "scan.pdf") as document:
    page = document[0]
    textpage = page.get_textpage_ocr(language="eng", tessdata=tessdata, dpi=150, full=True)
    contains(page.get_text(textpage=textpage), "pymupdf-offline-ocr")
contains(run("pdftotext", [pdf, "-"]), "poppler-text")
run("pdftoppm", ["-f", "1", "-singlefile", "-png", "-scale-to", "800", pdf, work / "poppler-render"])
with Image.open(work / "poppler-render.png") as rendered:
    if max(rendered.size) != 800:
        raise RuntimeError("Poppler did not render the expected dimensions")
proof["poppler-render"] = True
(work / "fixture.md").write_text("# " + marker + "\n\n| Article | Value |\n| --- | --- |\n| 1 | 42 |\n")
run("pandoc", [work / "fixture.md", "-o", work / "pandoc.docx"])
contains(run("pandoc", [work / "pandoc.docx", "-t", "plain"]), "pandoc-docx-roundtrip")
office = work / "office"
office.mkdir()
for ext in ("docx", "xlsx", "pptx"):
    run("libreoffice", ["-env:UserInstallation=" + (work / ("lo-" + ext)).as_uri(),
                        "--headless", "--convert-to", "pdf", "--outdir", office, work / ("fixture." + ext)])
    output = office / "fixture.pdf"
    contains(run("pdftotext", [output, "-"]), "libreoffice-" + ext + "-pdf")
    output.rename(office / (ext + ".pdf"))
run("libreoffice", ["-env:UserInstallation=" + (work / "lo-odt").as_uri(), "--headless",
                    "--convert-to", "odt", "--outdir", work, docx])
if not (work / "fixture.odt").is_file():
    raise RuntimeError("LibreOffice did not produce ODT")
run("qpdf", ["--empty", "--pages", pdf, pdf, "--", work / "merged.pdf"])
if len(pypdf.PdfReader(work / "merged.pdf").pages) != 2:
    raise RuntimeError("qpdf did not merge real pages")
run("qpdf", ["--check", work / "merged.pdf"])
proof["qpdf-merge"] = True
run("gs", ["-dSAFER", "-dBATCH", "-dNOPAUSE", "-sDEVICE=png16m", "-r72",
           "-dFirstPage=1", "-dLastPage=1", "-sOutputFile=" + str(work / "ghostscript.png"), pdf])
with Image.open(work / "ghostscript.png") as rendered:
    if not all(rendered.size):
        raise RuntimeError("Ghostscript render is empty")
proof["ghostscript-safer-render"] = True
run("ffmpeg", ["-v", "error", "-i", work / "scan.png", "-vf", "scale=900:250", "-frames:v", "1",
               "-threads", "1", work / "ffmpeg.png"])
with Image.open(work / "ffmpeg.png") as rendered:
    if rendered.size != (900, 250):
        raise RuntimeError("ffmpeg failed to transform the fixture")
proof["ffmpeg-image-conversion"] = True
run("convert", [work / "scan.png", "-resize", "900x250", work / "imagemagick.png"])
with Image.open(work / "imagemagick.png") as rendered:
    if rendered.size != (900, 250):
        raise RuntimeError("ImageMagick failed to resize the fixture")
proof["imagemagick-image-conversion"] = True
with zipfile.ZipFile(work / "fixture.zip", "w") as archive:
    archive.write(pdf, "nested/fixture.pdf")
    archive.write(docx, "nested/fixture.docx")
with gzip.open(work / "fixture.txt.gz", "wt") as compressed:
    compressed.write(marker)
for path, name in [(pdf, "rga-poppler-pdf"), (docx, "rga-pandoc-docx"),
                   (work / "fixture.zip", "rga-nested-zip"), (work / "fixture.txt.gz", "rga-gzip")]:
    contains(run(str(root / "rga/rga"), ["--rga-cache-path=" + str(work / "rga-cache"), marker, path]), name)
book = epub.EpubBook()
book.set_identifier("opulent-local-fixture")
book.set_title(marker)
book.set_language("en")
chapter = epub.EpubHtml(title=marker, file_name="fixture.xhtml", lang="en")
chapter.content = "<h1>" + marker + "</h1>"
book.add_item(chapter)
book.toc = (chapter,)
book.spine = [chapter]
book.add_item(epub.EpubNcx())
epub.write_epub(str(work / "fixture.epub"), book)
proof["ebooklib-epub-write"] = True
(work / "fixture.rtf").write_text("{\\rtf1\\ansi " + marker + "}")
(work / "fixture.csv").write_text("label,value\n" + marker + ",42\n")
frame = pandas.read_csv(work / "fixture.csv")
frame.to_parquet(work / "fixture.parquet", engine="pyarrow")
contains(pandas.read_parquet(work / "fixture.parquet").iloc[0]["label"], "pandas-pyarrow-roundtrip")
contains(BeautifulSoup("<p>" + marker + "</p>", "html.parser").get_text(), "beautifulsoup-html")
contains(etree.fromstring(("<p>" + marker + "</p>").encode()).text, "lxml-xml")
contains(ElementTree.fromstring("<p>" + marker + "</p>").text, "defusedxml-xml")
contains(yaml.safe_load("label: " + marker)["label"], "yaml-safe-load")
(work / "fixtures.json").write_text(json.dumps({"marker": marker, "proof": proof}))
print(json.dumps({"verified": True, "checks": proof,
                  "policy": "Generated local fixtures only; English OCR uses pinned prepared traineddata. No API credentials or hosted providers."}))
